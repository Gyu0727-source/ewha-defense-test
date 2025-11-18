#!/usr/bin/env python3
"""
이화방어기제검사 자동화 스크립트 v2.0

Google Sheets 응답 → 채점 엑셀 자동 입력

Requirements:
    pip install -r requirements.txt

사용 방법:
    python3 src/edmt_automation.py \\
        --sheet-url "YOUR_GOOGLE_SHEETS_URL" \\
        --template "이화방어검사채점표 수정.xlsx" \\
        --output-dir "./results"
"""

import argparse
import sys
import logging
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from datetime import datetime

try:
    import gspread
    from google.oauth2.service_account import Credentials
    from openpyxl import load_workbook
    import pandas as pd
except ImportError as e:
    print(f"❌ 필수 라이브러리가 설치되지 않았습니다: {e}")
    print("\n다음 명령어로 설치하세요:")
    print("pip install -r requirements.txt")
    sys.exit(1)


# 로깅 설정
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('edmt_automation.log', encoding='utf-8'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)


# 20개 방어기제
DEFENSE_MECHANISMS = [
    "허세", "반동형성", "동일시", "수동공격", "투사", "전치", "부정", "통제", "억제", "왜곡",
    "예견", "합리화", "해리", "신체화", "승화", "행동화", "이타주의", "퇴행", "유머", "회피"
]

# 규준표 매핑 (나이/성별)
NORM_TABLE_MAPPING = {
    (15, 29, "남"): "규준표11",
    (15, 29, "여"): "규준표12",
    (30, 39, "남"): "규준표21",
    (30, 39, "여"): "규준표22",
    (40, 49, "남"): "규준표31",
    (40, 49, "여"): "규준표32",
    (50, 150, "남"): "규준표41",  # 50세 이상
    (50, 150, "여"): "규준표42",
}


class EDMTAutomation:
    """이화방어기제검사 자동화 클래스"""

    def __init__(self, credentials_file: str = "credentials.json"):
        """
        초기화

        Args:
            credentials_file: Google API 인증 파일 경로
        """
        self.credentials_file = Path(credentials_file)
        self.client: Optional[gspread.Client] = None

    def authenticate(self) -> bool:
        """
        Google Sheets API 인증

        Returns:
            bool: 인증 성공 여부
        """
        try:
            if not self.credentials_file.exists():
                logger.error(f"인증 파일을 찾을 수 없습니다: {self.credentials_file}")
                logger.info("\nGoogle Cloud Console에서 서비스 계정 키를 생성하고")
                logger.info("credentials.json 파일로 저장하세요.")
                return False

            scopes = [
                'https://www.googleapis.com/auth/spreadsheets.readonly',
                'https://www.googleapis.com/auth/drive.readonly'
            ]

            creds = Credentials.from_service_account_file(
                str(self.credentials_file),
                scopes=scopes
            )
            self.client = gspread.authorize(creds)
            logger.info("✅ Google Sheets API 인증 성공")
            return True

        except Exception as e:
            logger.error(f"❌ 인증 실패: {e}", exc_info=True)
            return False

    def read_responses(self, sheet_url: str) -> List[Dict]:
        """
        Google Sheets에서 응답 읽기

        Args:
            sheet_url: Google Sheets URL

        Returns:
            list: 응답 데이터 리스트
        """
        try:
            logger.info(f"Google Sheets 열기: {sheet_url}")
            sheet = self.client.open_by_url(sheet_url)
            worksheet = sheet.get_worksheet(0)  # 첫 번째 시트

            # 모든 데이터 가져오기
            all_records = worksheet.get_all_records()

            logger.info(f"✅ {len(all_records)}개의 응답을 읽어왔습니다")
            return all_records

        except gspread.exceptions.APIError as e:
            logger.error(f"❌ Google Sheets API 오류: {e}", exc_info=True)
            return []
        except gspread.exceptions.SpreadsheetNotFound:
            logger.error("❌ 스프레드시트를 찾을 수 없습니다. URL을 확인하세요.")
            return []
        except Exception as e:
            logger.error(f"❌ Google Sheets 읽기 실패: {e}", exc_info=True)
            return []

    def parse_response(self, response_row: Dict) -> Optional[Dict]:
        """
        응답 데이터 파싱

        Args:
            response_row: 한 행의 응답 데이터 (dict)

        Returns:
            dict: 파싱된 데이터 또는 None (파싱 실패 시)
        """
        try:
            # 기본 정보 추출
            name = response_row.get('이름', '').strip()
            if not name:
                logger.warning("이름이 없는 응답을 건너뜁니다.")
                return None

            try:
                age = int(response_row.get('나이', 0))
                if age < 15 or age > 100:
                    logger.warning(f"{name}: 나이가 유효 범위를 벗어났습니다 ({age}세)")
                    return None
            except (ValueError, TypeError):
                logger.warning(f"{name}: 나이를 파싱할 수 없습니다")
                return None

            gender = response_row.get('성별', '').strip()
            if gender not in ['남', '여']:
                logger.warning(f"{name}: 성별이 유효하지 않습니다 ({gender})")
                return None

            # 200개 문항 응답 추출
            answers = []
            missing_count = 0

            for i in range(1, 201):
                question_key = f'{i}. '  # "1. 일생에 한 번뿐인..." 형식

                # 실제 키를 찾기
                full_key = None
                for key in response_row.keys():
                    if key.startswith(question_key):
                        full_key = key
                        break

                if full_key:
                    answer_text = str(response_row[full_key])
                    # "① 전혀 아니다" → 1
                    score = self._parse_answer_text(answer_text)
                    answers.append(score)

                    if score == 0:
                        missing_count += 1
                else:
                    logger.warning(f"{name}: 문항 {i}를 찾을 수 없습니다")
                    answers.append(0)
                    missing_count += 1

            # 무응답이 너무 많은 경우 경고
            if missing_count > 20:  # 10% 이상
                logger.warning(
                    f"{name}: 무응답이 {missing_count}개({missing_count/200*100:.1f}%)입니다"
                )

            return {
                'name': name,
                'age': age,
                'gender': gender,
                'answers': answers,
                'missing_count': missing_count
            }

        except Exception as e:
            logger.error(f"응답 파싱 실패: {e}", exc_info=True)
            return None

    def _parse_answer_text(self, answer_text: str) -> int:
        """
        응답 텍스트를 점수로 변환

        Args:
            answer_text: 응답 텍스트 (예: "① 전혀 아니다")

        Returns:
            int: 점수 (1-5, 무응답은 0)
        """
        if '①' in answer_text or '1' in answer_text:
            return 1
        elif '②' in answer_text or '2' in answer_text:
            return 2
        elif '③' in answer_text or '3' in answer_text:
            return 3
        elif '④' in answer_text or '4' in answer_text:
            return 4
        elif '⑤' in answer_text or '5' in answer_text:
            return 5
        else:
            logger.debug(f"알 수 없는 응답 형식: {answer_text}")
            return 0  # 무응답

    def calculate_raw_scores(self, answers: List[int]) -> Dict[str, int]:
        """
        원점수 계산

        Args:
            answers: 200개 문항 응답 (list)

        Returns:
            dict: 각 방어기제별 원점수
        """
        scores = {}

        for mech_idx, mechanism in enumerate(DEFENSE_MECHANISMS):
            # 해당 방어기제의 문항 번호들 (1-based)
            # 예: 허세 = 1, 21, 41, 61, 81, 101, 121, 141, 161, 181
            item_numbers = [mech_idx + 1 + (i * 20) for i in range(10)]

            # 원점수 합산
            total_score = sum(answers[item_num - 1] for item_num in item_numbers)
            scores[mechanism] = total_score

        return scores

    def get_norm_table_code(self, age: int, gender: str) -> str:
        """
        나이와 성별에 따른 규준표 코드 결정

        Args:
            age: 나이
            gender: 성별 ("남" or "여")

        Returns:
            str: 규준표 이름 (예: "규준표11")
        """
        for (age_min, age_max, norm_gender), table_name in NORM_TABLE_MAPPING.items():
            if age_min <= age <= age_max and gender == norm_gender:
                return table_name

        # 기본값 (예외 처리)
        logger.warning(f"규준표를 찾을 수 없습니다. 나이: {age}, 성별: {gender}. 기본값 사용")
        return "규준표11" if gender == "남" else "규준표12"

    def convert_to_sten(
        self,
        raw_score: int,
        norm_table_ws,
        mechanism_col: int
    ) -> int:
        """
        원점수를 스텐점수로 변환

        Args:
            raw_score: 원점수
            norm_table_ws: 규준표 워크시트
            mechanism_col: 방어기제 열 인덱스 (1-based)

        Returns:
            int: 스텐점수 (1-10)
        """
        try:
            # 규준표에서 스텐점수 찾기
            # Row 2-11이 스텐점수 1-10에 해당
            for sten_score in range(1, 11):
                row_idx = sten_score + 1  # 엑셀 행 번호 (2-11)
                threshold = norm_table_ws.cell(row=row_idx, column=mechanism_col).value

                if threshold is None:
                    logger.warning(f"규준표 데이터 누락: row={row_idx}, col={mechanism_col}")
                    continue

                if raw_score <= int(threshold):
                    return sten_score

            return 10  # 최대값

        except Exception as e:
            logger.error(f"스텐점수 변환 실패: {e}", exc_info=True)
            return 5  # 기본값 (중간)

    def write_to_excel(
        self,
        template_path: Path,
        output_path: Path,
        parsed_data: Dict,
        raw_scores: Dict[str, int]
    ) -> bool:
        """
        채점 엑셀에 데이터 입력

        Args:
            template_path: 템플릿 엑셀 경로
            output_path: 출력 엑셀 경로
            parsed_data: 파싱된 응답 데이터
            raw_scores: 원점수 딕셔너리

        Returns:
            bool: 성공 여부
        """
        try:
            # 엑셀 파일 로드
            wb = load_workbook(template_path)

            # 응답지_1 시트에 응답 입력
            if '응답지_1' not in wb.sheetnames:
                logger.error("템플릿에 '응답지_1' 시트가 없습니다")
                return False

            response_sheet = wb['응답지_1']

            # 기본 정보 입력 (행 4)
            response_sheet['D4'] = parsed_data['name']  # 이름
            response_sheet['F4'] = parsed_data['age']    # 나이
            response_sheet['H4'] = parsed_data['gender'] # 성별

            # 200개 응답 입력
            answers = parsed_data['answers']
            for i, answer in enumerate(answers):
                # 20문항씩 나눠서 입력 (행 7, 10, 13, 16, 19, 22, 25, 28, 31, 34)
                section = i // 20
                position_in_section = i % 20

                # 행 번호 계산
                base_row = 7 + (section * 3)
                # 열 번호 계산 (D열부터 시작, 20개)
                col = 4 + position_in_section  # D=4

                response_sheet.cell(row=base_row, column=col, value=answer)

            # 저장
            wb.save(output_path)
            logger.info(f"✅ 엑셀 파일 저장 완료: {output_path}")
            return True

        except FileNotFoundError:
            logger.error(f"템플릿 파일을 찾을 수 없습니다: {template_path}")
            return False
        except Exception as e:
            logger.error(f"엑셀 저장 실패: {e}", exc_info=True)
            return False

    def process_all_responses(
        self,
        sheet_url: str,
        template_path: str,
        output_dir: str = "./results"
    ) -> Tuple[int, int]:
        """
        모든 응답 자동 처리

        Args:
            sheet_url: Google Sheets URL
            template_path: 템플릿 엑셀 경로
            output_dir: 출력 디렉토리

        Returns:
            tuple: (성공 개수, 실패 개수)
        """
        # 인증
        if not self.authenticate():
            return 0, 0

        # 응답 읽기
        responses = self.read_responses(sheet_url)

        if not responses:
            logger.error("❌ 처리할 응답이 없습니다")
            return 0, 0

        # 출력 디렉토리 생성
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)

        template = Path(template_path)
        if not template.exists():
            logger.error(f"템플릿 파일을 찾을 수 없습니다: {template_path}")
            return 0, 0

        # 각 응답 처리
        success_count = 0
        fail_count = 0

        for idx, response_row in enumerate(responses, 1):
            logger.info(f"\n{'=' * 60}")
            logger.info(f"응답 {idx}/{len(responses)} 처리 중...")
            logger.info(f"{'=' * 60}")

            try:
                # 응답 파싱
                parsed = self.parse_response(response_row)
                if parsed is None:
                    logger.warning(f"응답 {idx} 파싱 실패, 건너뜁니다")
                    fail_count += 1
                    continue

                logger.info(
                    f"📝 이름: {parsed['name']}, 나이: {parsed['age']}, "
                    f"성별: {parsed['gender']}, 무응답: {parsed['missing_count']}개"
                )

                # 원점수 계산
                raw_scores = self.calculate_raw_scores(parsed['answers'])
                logger.info("📊 원점수 계산 완료")

                # 엑셀에 저장
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_filename = f"{parsed['name']}_{parsed['age']}세_{parsed['gender']}_{timestamp}.xlsx"
                output_file = output_path / output_filename

                if self.write_to_excel(template, output_file, parsed, raw_scores):
                    success_count += 1
                else:
                    fail_count += 1

            except Exception as e:
                logger.error(f"응답 {idx} 처리 중 오류 발생: {e}", exc_info=True)
                fail_count += 1

        logger.info(f"\n{'=' * 60}")
        logger.info(f"✅ 전체 처리 완료!")
        logger.info(f"성공: {success_count}개, 실패: {fail_count}개")
        logger.info(f"📁 결과 파일 위치: {output_dir}/")
        logger.info(f"{'=' * 60}")

        return success_count, fail_count


def main():
    """메인 함수"""
    parser = argparse.ArgumentParser(
        description='이화방어기제검사 자동화 v2.0',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
예제:
  python3 src/edmt_automation.py \\
    --sheet-url "https://docs.google.com/spreadsheets/d/..." \\
    --template "이화방어검사채점표 수정.xlsx" \\
    --output-dir "./results"
        """
    )

    parser.add_argument(
        '--sheet-url',
        required=True,
        help='Google Sheets URL'
    )
    parser.add_argument(
        '--template',
        required=True,
        help='채점 템플릿 엑셀 파일 경로'
    )
    parser.add_argument(
        '--output-dir',
        default='./results',
        help='출력 디렉토리 (기본값: ./results)'
    )
    parser.add_argument(
        '--credentials',
        default='credentials.json',
        help='Google API 인증 파일 (기본값: credentials.json)'
    )
    parser.add_argument(
        '--log-level',
        choices=['DEBUG', 'INFO', 'WARNING', 'ERROR'],
        default='INFO',
        help='로그 레벨 (기본값: INFO)'
    )

    args = parser.parse_args()

    # 로그 레벨 설정
    logger.setLevel(getattr(logging, args.log_level))

    # 자동화 실행
    automation = EDMTAutomation(credentials_file=args.credentials)
    success, fail = automation.process_all_responses(
        sheet_url=args.sheet_url,
        template_path=args.template,
        output_dir=args.output_dir
    )

    # 종료 코드
    sys.exit(0 if fail == 0 else 1)


if __name__ == '__main__':
    main()
